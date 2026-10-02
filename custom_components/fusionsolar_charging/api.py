"""FusionSolar charge-record client with month cache, name mapping and exports."""
from __future__ import annotations

import asyncio
import calendar
import csv
import io
import logging
import threading
from datetime import date, datetime, time, timedelta
from zoneinfo import ZoneInfo

import requests
from fusion_solar_py.client import FusionSolarClient
from fusion_solar_py.exceptions import (
    AuthenticationException,
    CaptchaRequiredException,
)

DOMAIN = "fusionsolar_charging"
PATH = "/rest/neteco/web/homemgr/v2/charger/list-charge-record"
TTL = 300  # seconds before the running month is fetched again
_LOGGER = logging.getLogger(__name__)


class LoginError(Exception):
    """Error with a kind (auth, captcha, network, api) and a readable message."""

    def __init__(self, kind: str, message: str):
        super().__init__(message)
        self.kind, self.message = kind, message


def explain(err: Exception) -> LoginError:
    if isinstance(err, LoginError):
        return err
    if isinstance(err, CaptchaRequiredException):
        return LoginError("captcha", "FusionSolar asks for a captcha at login. "
                          "Log in once in a browser on the FusionSolar website, "
                          "then try again later.")
    if isinstance(err, AuthenticationException):
        return LoginError("auth", "FusionSolar rejected the login. Check the "
                          "username, password and subdomain.")
    if isinstance(err, (requests.ConnectionError, requests.Timeout)):
        return LoginError("network", "Could not reach FusionSolar. Try again "
                          "in a minute.")
    return LoginError("api", f"FusionSolar returned an error: {err}")


def login_client(username, password, subdomain):
    return FusionSolarClient(username, password, huawei_subdomain=subdomain)


def _month_end(m: str) -> date:
    y, mo = map(int, m.split("-"))
    return date(y, mo, calendar.monthrange(y, mo)[1])


def _months(s: date, e: date) -> list[str]:
    out, d = [], s.replace(day=1)
    while d <= e:
        out.append(d.strftime("%Y-%m"))
        d = (d + timedelta(days=32)).replace(day=1)
    return out


class Api:
    def __init__(self, hass, entry, store):
        d = entry.data
        self.hass, self.entry, self.store = hass, entry, store
        self.username, self.password = d["username"], d["password"]
        self.sub, self.dn_id = d["subdomain"], int(d["dn_id"])
        self.tz_name = hass.config.time_zone
        self.tz = ZoneInfo(self.tz_name)
        self.client = None
        self.lock = threading.Lock()
        self.alock = asyncio.Lock()
        self.data = {"months": {}, "names": {}}
        self.blocked = None  # (LoginError, until or None)

    async def async_load(self):
        saved = await self.store.async_load()
        if saved:
            self.data["months"] = saved.get("months", {})
            self.data["names"] = saved.get("names", {})

    # ---- blocking part (executor) ----
    def _epoch(self, day: date, end=False) -> int:
        t = time(23, 59, 59) if end else time(0, 0, 0)
        return int(datetime.combine(day, t, tzinfo=self.tz).timestamp())

    def _pages(self, start: date, end: date):
        offset = int(datetime.now(self.tz).utcoffset().total_seconds() // 60)
        rows, page = [], 1
        while True:
            r = self.client._session.post(
                f"https://{self.sub}.fusionsolar.huawei.com{PATH}",
                json={"startTime": str(self._epoch(start)),
                      "endTime": str(self._epoch(end, True)),
                      "pageNo": page, "pageSize": 1000,
                      "dnId": self.dn_id, "timeZoneId": self.tz_name},
                headers={"X-Timezone-Offset": str(offset)}, timeout=30)
            d = r.json()
            if d.get("code") != 0:
                raise ValueError(d.get("description") or "unknown API error")
            recs = (d.get("data") or {}).get("records") or []
            rows += recs
            if len(recs) < 1000:
                return rows
            page += 1

    def _fetch(self, start: date, end: date):
        with self.lock:
            for attempt in (0, 1):
                try:
                    if self.client is None or attempt:
                        self.client = login_client(
                            self.username, self.password, self.sub)
                    return self._pages(start, end)
                except (AuthenticationException, CaptchaRequiredException):
                    raise
                except Exception:  # noqa: BLE001  (session may have expired)
                    if attempt:
                        raise

    def _session(self, r) -> dict:
        s = datetime.fromtimestamp(int(r["startTime"]), self.tz)
        e = datetime.fromtimestamp(int(r["stopTime"]), self.tz)
        return {"id": r.get("orderNumber"), "day": s.strftime("%Y-%m-%d"),
                "start": s.strftime("%H:%M"), "end": e.strftime("%H:%M"),
                "endDay": e.strftime("%Y-%m-%d"),
                "user": r.get("accountId") or "unknown",
                "kwh": round(float(r.get("totalPower") or 0), 3),
                "minutes": r.get("totalTime")}

    # ---- async part ----
    async def _refresh_months(self, need: list[str], now: datetime):
        if self.blocked and (self.blocked[1] is None or now < self.blocked[1]):
            raise self.blocked[0]
        first = date.fromisoformat(min(need) + "-01")
        last = _month_end(max(need))
        try:
            raw = await self.hass.async_add_executor_job(self._fetch, first, last)
        except Exception as err:  # noqa: BLE001
            le = explain(err)
            if le.kind == "auth":
                self.blocked = (le, None)  # don't hammer a wrong password
                self.entry.async_start_reauth(self.hass)
            elif le.kind == "captcha":
                self.blocked = (le, now + timedelta(minutes=30))
            raise le from err
        self.blocked = None
        by = {m: [] for m in need}
        for r in raw:
            s = self._session(r)
            if s["day"][:7] in by:
                by[s["day"][:7]].append(s)
        today = now.date()
        for m in need:
            by[m].sort(key=lambda x: (x["day"], x["start"]))
            self.data["months"][m] = {
                "fetched": now.isoformat(), "complete": today > _month_end(m),
                "sessions": by[m]}
        await self.store.async_save(self.data)

    async def async_sessions(self, start: str, end: str, refresh=False):
        s, e = date.fromisoformat(start), date.fromisoformat(end)
        months = _months(s, e)
        async with self.alock:
            now = datetime.now(self.tz)
            need = []
            for m in months:
                c = self.data["months"].get(m)
                if c is None or refresh:
                    need.append(m)
                elif not c["complete"] and (
                        now - datetime.fromisoformat(c["fetched"])
                ).total_seconds() > TTL:
                    need.append(m)
            warning = None
            if need:
                try:
                    await self._refresh_months(need, now)
                except LoginError as le:
                    if any(m not in self.data["months"] for m in months):
                        raise
                    warning = le.message  # serve saved data instead
        names = self.data["names"]
        out = []
        for m in months:
            for x in (self.data["months"].get(m) or {}).get("sessions", []):
                if start <= x["day"] <= end:
                    out.append(dict(x, name=names.get(x["user"]) or x["user"]))
        out.sort(key=lambda x: (x["day"], x["start"]))
        fetched = [self.data["months"][m]["fetched"] for m in months
                   if m in self.data["months"]]
        return {"sessions": out, "names": dict(names),
                "updated": max(fetched) if fetched else None,
                "warning": warning}

    async def async_set_names(self, names: dict) -> dict:
        for k, v in names.items():
            k, v = str(k).strip(), str(v or "").strip()
            if v:
                self.data["names"][k] = v
            else:
                self.data["names"].pop(k, None)
        await self.store.async_save(self.data)
        return dict(self.data["names"])


# ---- exports ----
def summarize(rows):
    agg = {}
    for r in rows:
        a = agg.setdefault(r["user"], [r["name"], 0, 0.0])
        a[1] += 1
        a[2] += r["kwh"]
    return sorted(((n, u, c, k) for u, (n, c, k) in agg.items()),
                  key=lambda x: -x[3])


def _n(x):
    return f"{x:.3f}".replace(".", ",")  # decimal comma for Belgian Excel


def make_csv(rows) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Date", "Start", "End", "End date", "User", "Account ID",
                "Duration (min)", "kWh", "Session ID"])
    for r in rows:
        w.writerow([r["day"], r["start"], r["end"], r["endDay"], r["name"],
                    r["user"], r["minutes"], _n(r["kwh"]), r["id"]])
    w.writerow(["TOTAL", "", "", "", "", "", "",
                _n(sum(r["kwh"] for r in rows)), f"{len(rows)} sessions"])
    w.writerow([])
    w.writerow(["Totals per user"])
    w.writerow(["User", "Account ID", "Sessions", "kWh"])
    for n, u, c, k in summarize(rows):
        w.writerow([n, u, c, _n(k)])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def make_xlsx(rows, period: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Sessions"
    ws.append(["Date", "Start", "End", "End date", "User", "Account ID",
               "Duration (min)", "kWh", "Session ID"])
    for r in rows:
        ws.append([date.fromisoformat(r["day"]), r["start"], r["end"],
                   date.fromisoformat(r["endDay"]), r["name"], r["user"],
                   r["minutes"], r["kwh"], r["id"]])
    n = len(rows) + 1
    ws.append(["TOTAL", "", "", "", "", "", "", f"=SUM(H2:H{n})" if rows else 0,
               f"{len(rows)} sessions"])
    for row in ws.iter_rows(min_row=2, max_row=n):
        row[0].number_format = row[3].number_format = "yyyy-mm-dd"
        row[7].number_format = "0.000"
    ws.cell(n + 1, 8).number_format = "0.000"
    for c in ws[1] + ws[n + 1]:
        c.font = Font(bold=True)
    ws.freeze_panes = "A2"
    for col, wd in zip("ABCDEFGHI", (12, 8, 8, 12, 22, 16, 15, 10, 34)):
        ws.column_dimensions[col].width = wd

    s2 = wb.create_sheet("Per user")
    s2.append([f"Period {period}"])
    s2.append(["User", "Account ID", "Sessions", "kWh"])
    summ = summarize(rows)
    for item in summ:
        s2.append(list(item))
    last = len(summ) + 2
    s2.append(["TOTAL", "", f"=SUM(C3:C{last})" if summ else 0,
               f"=SUM(D3:D{last})" if summ else 0])
    for row in s2.iter_rows(min_row=3, max_row=last + 1, min_col=4, max_col=4):
        row[0].number_format = "0.000"
    for c in s2[2] + s2[last + 1]:
        c.font = Font(bold=True)
    s2["A1"].font = Font(bold=True)
    for col, wd in zip("ABCD", (24, 16, 10, 10)):
        s2.column_dimensions[col].width = wd
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
